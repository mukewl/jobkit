"""Generate the two workbooks the process runs on.

    python trackers/make_trackers.py            # creates both, refuses to overwrite
    python trackers/make_trackers.py --force    # overwrite

  Job_Tracker.xlsx   one row per role: tier, language risk, link, status
  Contacts.xlsx      one row per person: why them, both drafted messages, follow-up date

Both are gitignored. They are your data, not the repo's.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
ARIAL = "Arial"

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TIER = {
    "A": PatternFill("solid", fgColor="C6EFCE"),
    "B": PatternFill("solid", fgColor="FFEB9C"),
    "C": PatternFill("solid", fgColor="FCE4D6"),
    "X": PatternFill("solid", fgColor="E7E6E6"),
}
MUTED = Font(name=ARIAL, size=9, italic=True, color="808080")

EXAMPLE_NOTE = "Example row - delete it"

CONNECT_NOTE = (
    "Hi [Name] - I've just applied for the [Role] role at [Company]. "
    "[One line on who you are.] Would like to connect."
)

FOLLOW_MSG = (
    "Hi [Name], thanks for connecting.\n\n"
    "I've just applied for the [Role] role at [Company], and I'd rather ask someone "
    "inside than work it out from the job ad.\n\n"
    "[Two sentences on you, tied to something specific about them. No metrics.]\n\n"
    "I'd like to know [what the team is actually like day to day], and "
    "[something only they could answer].\n\n"
    "Would you have 15 minutes for a call in the next week or two? Happy to fit "
    "around whenever suits you.\n\n"
    "Thanks,\n[Your name]"
)


def header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def notes(ws, row, lines):
    for i, text in enumerate(lines):
        ws.cell(row + i, 1, text).font = MUTED


def style_row(ws, row, ncols, wrap):
    for col in range(1, ncols + 1):
        cell = ws.cell(row, col)
        cell.border = BORDER
        cell.font = Font(name=ARIAL, size=10)
        cell.alignment = Alignment(vertical="top", wrap_text=(col in wrap))


# ---------------------------------------------------------------- job tracker

def build_job_tracker(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"

    cols = ["#", "Tier", "Company", "Role", "Location", "Work auth", "Score",
            "Source", "Job link", "Status", "Date applied", "CV file", "Notes"]
    header(ws, cols)

    ws.append([1, "A", "Northwind Labs", "Growth Marketing Manager", "Paris, France",
               "OK", 88, "csv_import", "https://example.com/jobs/1001234",
               "Not started", None, None, EXAMPLE_NOTE])
    style_row(ws, 2, len(cols), wrap=(4, 13))
    ws.cell(2, 2).fill = TIER["A"]
    ws.cell(2, 2).alignment = Alignment(horizontal="center", vertical="top")

    widths(ws, {1: 5, 2: 6, 3: 24, 4: 42, 5: 20, 6: 12, 7: 7, 8: 14,
                9: 34, 10: 14, 11: 13, 12: 26, 13: 34})
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:M{}".format(ws.max_row)

    notes(ws, 5, [
        "Fill in only: Status, Date applied, CV file, Notes. The rest is written for you.",
        "Tier   A = apply now | B = stretch or adjacent | C = check the language bar first | X = you cannot take it",
        "Status values: Not started / Applied / Interview / Rejected / Offer",
    ])

    s = wb.create_sheet("Summary")
    s["A1"] = "Applications"
    s["A1"].font = Font(name=ARIAL, bold=True, size=14)
    last = 500
    rows = [
        ("Total roles", '=COUNTA(Applications!$C$2:$C${})'.format(last)),
        ("Tier A", '=COUNTIF(Applications!$B$2:$B${},"A")'.format(last)),
        ("Tier B", '=COUNTIF(Applications!$B$2:$B${},"B")'.format(last)),
        ("Tier C", '=COUNTIF(Applications!$B$2:$B${},"C")'.format(last)),
        ("Applied", '=COUNTIF(Applications!$J$2:$J${},"Applied")'.format(last)),
        ("Interviewing", '=COUNTIF(Applications!$J$2:$J${},"Interview")'.format(last)),
        ("Rejected", '=COUNTIF(Applications!$J$2:$J${},"Rejected")'.format(last)),
        ("Still to do", '=COUNTIF(Applications!$J$2:$J${},"Not started")'.format(last)),
    ]
    for i, (label, formula) in enumerate(rows, start=3):
        s.cell(i, 1, label).font = Font(name=ARIAL, size=10)
        s.cell(i, 2, formula).font = Font(name=ARIAL, size=10, bold=True)
    s.column_dimensions["A"].width = 18
    s.column_dimensions["B"].width = 12
    notes(s, 13, ["Counts recalculate when Excel opens the file."])

    wb.save(path)


# ---------------------------------------------------------------- contacts

def build_contacts(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"

    cols = ["Company", "Role applied", "Priority", "Name", "Title", "Why them",
            "LinkedIn", "Connection note (<300 char)", "Message once connected",
            "Status", "Date sent", "Follow-up due", "Notes"]
    header(ws, cols)

    ws.append([
        "Northwind Labs", "Growth Marketing Manager", 1, "A. Example", "Head of Growth",
        "Owns the team this role sits in", "linkedin.com/in/example",
        CONNECT_NOTE, FOLLOW_MSG, "Not sent", None, None, EXAMPLE_NOTE,
    ])
    style_row(ws, 2, len(cols), wrap=(5, 6, 8, 9, 13))
    ws.cell(2, 3).fill = TIER["A"]
    ws.cell(2, 3).alignment = Alignment(horizontal="center", vertical="top")
    ws.row_dimensions[2].height = 200

    # Follow-up due fills itself once a request has gone out and no reply has landed.
    for r in range(2, 400):
        f = ('=IF(AND(K{r}<>"",OR(J{r}="Request sent",J{r}="Connected")),K{r}+7,"")'
             .format(r=r))
        cell = ws.cell(r, 12, f)
        cell.font = Font(name=ARIAL, size=10)
        cell.number_format = "yyyy-mm-dd"
        cell.alignment = Alignment(vertical="top")

    widths(ws, {1: 16, 2: 26, 3: 8, 4: 18, 5: 32, 6: 40, 7: 42, 8: 46,
                9: 72, 10: 13, 11: 12, 12: 13, 13: 26})
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = "A1:M2"

    notes(ws, 403, [
        "Priority   1-3 = message these | 4-5 = optional | 9 = found by search but not relevant",
        "Fill in only: Status, Date sent, Notes. Follow-up due calculates itself.",
        "Status values: Not sent / Request sent / Connected / Replied / No reply",
        "Message style: ask, do not pitch. No metrics in outreach - they belong on the CV.",
    ])

    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="overwrite existing workbooks")
    args = ap.parse_args()

    targets = [
        (ROOT / "Job_Tracker.xlsx", build_job_tracker),
        (ROOT / "Contacts.xlsx", build_contacts),
    ]

    existing = [p.name for p, _ in targets if p.exists()]
    if existing and not args.force:
        print("refusing to overwrite: {}".format(", ".join(existing)))
        print("  pass --force if you really mean it (this deletes your tracked data)")
        return 1

    for path, builder in targets:
        builder(path)
        print("wrote {}".format(path))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
